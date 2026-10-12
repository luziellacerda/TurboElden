from pathlib import Path
from datetime import datetime, timezone, timedelta
import hashlib
import json
import re
import zipfile
import unicodedata

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

import argparse
parser=argparse.ArgumentParser(description="Render the public handoff chapters; never reads phones, APKs or credentials")
parser.add_argument('--input',type=Path,required=True)
parser.add_argument('--docx',type=Path,required=True)
parser.add_argument('--markdown',type=Path,required=True)
args=parser.parse_args()
data=json.loads((args.input/'chapters.json').read_text(encoding='utf-8'))
groups=list(data['groups'])
review=args.input/'consumer-review.json'
if review.exists():groups.insert(1,json.loads(review.read_text(encoding='utf-8')))
audit=args.input/'audit-chapters.json'
if audit.exists():groups.append(json.loads(audit.read_text(encoding='utf-8')))
all_sources=data['sources']
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.8)
sec.bottom_margin = Inches(.8)
sec.left_margin = Inches(.8)
sec.right_margin = Inches(.8)
sec.header_distance = Inches(.32)
sec.footer_distance = Inches(.35)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number','Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.13
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name,size in [('Title',26),('Subtitle',13),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(14 if name.startswith('Heading') else 0)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True
for name in ['List Bullet','List Number']:
    doc.styles[name].font.size = Pt(10.5)
    doc.styles[name].paragraph_format.space_after = Pt(4)
header = sec.header.paragraphs[0]
header.text = 'TurboStations R134   Dossiê técnico e handoff'
header.style = doc.styles['Normal']
header.runs[0].font.size = Pt(8.5)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Página ').font.size = Pt(9)
def field(p, value):
    run = p.add_run()
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'),value)
    run._r.addnext(f)
field(footer,'PAGE')
footer.add_run(' de ').font.size = Pt(9)
field(footer,'NUMPAGES')
settings = doc.settings.element
update = OxmlElement('w:updateFields')
update.set(qn('w:val'),'true')
settings.append(update)

def wrap(value):
    # Visual break opportunities preserve readable long technical paths/hashes.
    value = str(value)
    value = re.sub(r'\b[0-9a-f]{40,64}\b', lambda m:'\u200b'.join(m.group(0)[i:i+16] for i in range(0,len(m.group(0)),16)), value)
    return re.sub(r'(?<=[/\\_])(?=\S)', '\u200b', value)

def hyperlink(p, text, url=None, anchor=None, size=None):
    h = OxmlElement('w:hyperlink')
    if url:
        h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    if anchor:
        h.set(qn('w:anchor'),anchor)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color');color.set(qn('w:val'),'174A70');pr.append(color)
    if size is not None:
        font_size=OxmlElement('w:sz');font_size.set(qn('w:val'),str(int(size*2)));pr.append(font_size)
    r.append(pr)
    t = OxmlElement('w:t');t.text = wrap(text);r.append(t)
    h.append(r);p._p.append(h)

def rich(p, text):
    text = str(text)
    text = re.sub(r',(?=[A-Za-z_]|\d+x\d+)', ', ', text)
    words='documentado|admite|cada|convite|conserva|limita|carrega|autorizado|tentativas|indisponibilidade|texto|Falha|Vencer|cobre|que|respostas|Recusa|qualquer|crescimento|acima|tem|aceita|reserva|sucesso'
    text=re.sub(r'\b('+words+r')(?=\d)',r'\1 ',text)
    for old,new in [('PPSSPP1.20.4','PPSSPP 1.20.4'),('YabaSanshiro1.20.46','YabaSanshiro 1.20.46'),('XenDroid0b11201','XenDroid 0b11201'),('Cemu0.5.2','Cemu 0.5.2'),('MAME0.289','MAME 0.289'),('Dolphin2609-7','Dolphin 2609-7'),('aPS3e2.43','aPS3e 2.43'),('Android9','Android 9'),('Android12','Android 12'),('Android30','Android API 30'),('API28','API 28'),('API31','API 31'),('Config18campos','Config 18 campos'),('Config11campos','Config 11 campos')]:
        text=text.replace(old,new)
    pattern = re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*')
    cursor = 0
    for m in pattern.finditer(text):
        p.add_run(wrap(text[cursor:m.start()]))
        if m.group(1): hyperlink(p,m.group(1),m.group(2))
        elif m.group(3):
            r = p.add_run(wrap(m.group(3)));r.font.name='Consolas';r.font.size=Pt(10)
        else: p.add_run(wrap(m.group(4))).bold=True
        cursor=m.end()
    p.add_run(wrap(text[cursor:]))
    return p

def prose(text, style=None):
    p=doc.add_paragraph(style=style)
    return rich(p,text)

def clean_heading(text):
    return re.sub(r'\s+',' ', re.sub(r'[^\w\sÀ-ÿ]',' ',str(text))).strip()

bookmark_count = 0
def heading(text, level=2, anchor=None):
    global bookmark_count
    p=doc.add_heading(clean_heading(text),level)
    if anchor:
        bookmark_count += 1
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_count));start.set(qn('w:name'),anchor)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_count))
        p._p.insert(1,start);p._p.append(end)
    return p

def table(headers, rows, widths=None, font=9.4):
    n=len(headers)
    if widths is None:
        widths=([2.05,4.85] if n==2 else [1.55,2.05,3.30] if n==3 else [1.05,2.00,1.25,2.60] if n==4 else [6.9/n]*n)
    t=doc.add_table(rows=1,cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for name,v in [('top',85),('bottom',85),('left',105),('right',105)]:
        e=OxmlElement('w:'+name);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');margins.append(e)
    t._tbl.tblPr.append(margins)
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    data=[headers]+rows
    for idx,values in enumerate(data):
        row=t.rows[0] if idx==0 else t.add_row()
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill='293846' if idx==0 else 'F3F5F7' if idx%2==0 else 'FFFFFF'
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
            rich(p,value)
            for r in p.runs:
                r.font.name='Arial';r.font.size=Pt(font)
                r.font.color.rgb=RGBColor(255,255,255) if idx==0 else RGBColor(0,0,0)
                r.bold=idx==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


doc.add_paragraph('Dossiê completo do TurboStations e handoff de manutenção',style='Title')
doc.add_paragraph('Aplicativo R134 auditoria online servidor e entrega ao consumidor',style='Subtitle')
prose('Referência executável '+data['referenceCommit']+'. APK instalado R134 SHA256 '+data['apkSha256']+'.')
prose('O proprietário relatou carregamento de partida com três jogadores e pedido para fechar o aplicativo ao sair. Esta edição consolida os capítulos do aplicativo completo com a auditoria atual e as correções em fonte, sem usar USB. A homologação física da correção permanece pendente.')
prose('Consultar os inventários estruturados e os relatórios de testes junto do documento. Fontes e receitas públicas têm dependências privadas indicadas por hash; o repositório público não contém APK, firmware, BIOS, jogos ou chaves.')
doc.add_page_break();heading('Sumário de consulta',1)
for gi,g in enumerate(groups,1):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(8);hyperlink(p,f"{gi}  {g['title']}",anchor=f'group{gi}',size=10)
    for si,s in enumerate(g['sections'],1):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
        hyperlink(p,f"{gi}.{si}  {clean_heading(s['title'])}",anchor=f'g{gi}s{si}',size=9.5)
for gi,g in enumerate(groups,1):
    doc.add_page_break();heading(g['title'],1,f'group{gi}')
    for si,s in enumerate(g['sections'],1):
        heading(f"{gi}.{si} {s['title']}",2,f'g{gi}s{si}')
        for text in s.get('paragraphs',[]):prose(text)
        for t in s.get('tables',[]):table(t['headers'],t['rows'])
        for b in s.get('bullets',[]):prose(b,'List Bullet')
doc.add_page_break();heading('Inventário de bibliotecas e DEX',1)
for name in ('dex-entries.json','dex-assets.json','native-libraries.json'):
    rows=json.loads((args.input/name).read_text(encoding='utf-8'))
    table(['Entrada','Bytes','SHA256'],[[r['path'],str(r['bytes']),r['sha256']] for r in rows],widths=[3.2,.8,2.9],font=8.7)
heading('Fontes produtivas congeladas da R134',1)
frozen=json.loads((args.input/'productive-sources.json').read_text(encoding='utf-8'))
table(['Caminho na integração R134','SHA256'],list(map(list,frozen.items())),widths=[3.55,3.35],font=9)
doc.add_page_break();heading('Índice das fontes consultadas',1)
for s in all_sources:
    p=doc.add_paragraph();p.add_run(str(s.get('id',''))+'  ').bold=True;rich(p,s.get('purpose') or s.get('description') or '')
    path=s.get('path') or s.get('localPath');url=s.get('url')
    p.paragraph_format.keep_with_next=bool(path or url)
    if path:
        p=prose('Caminho da fonte '+str(path));p.paragraph_format.keep_with_next=bool(url);p.paragraph_format.keep_together=True
    if url:
        p=doc.add_paragraph();hyperlink(p,'Fonte '+str(s.get('id','')),url)
doc.core_properties.title='Dossiê completo do TurboStations e handoff de manutenção'
doc.core_properties.subject='Aplicativo R134 auditoria online servidor e entrega ao consumidor'
doc.core_properties.author='Codex'
args.docx.parent.mkdir(parents=True,exist_ok=True);doc.save(args.docx)

def cell(x):return str(x).replace('|',chr(92)+'|').replace('\n','<br>')
lines=['# Dossiê completo do TurboStations e handoff de manutenção','',
       'Aplicativo R134, auditoria online, servidor e entrega ao consumidor.','',
       '**Fonte executável:** '+data['referenceCommit']+'. **APK R134:** '+data['apkSha256']+'.','',
       'O relato de carregamento com três jogadores e de pedido para fechar o app está registrado na auditoria atual. Correções em fonte não equivalem a uma nova instalação homologada.','']
for gi,g in enumerate(groups,1):
    lines+=['## '+str(gi)+' '+g['title'],'']
    for si,s in enumerate(g['sections'],1):
        lines+=['### '+str(gi)+'.'+str(si)+' '+s['title'],'']
        for text in s.get('paragraphs',[]):lines+=[text,'']
        for t in s.get('tables',[]):
            lines+=['| '+' | '.join(map(cell,t['headers']))+' |','| '+' | '.join(['---']*len(t['headers']))+' |']
            lines+=['| '+' | '.join(map(cell,row))+' |' for row in t['rows']];lines+=['']
        lines+=['- '+str(b) for b in s.get('bullets',[])];lines+=['']
lines+=['## Índices completos','',
       '- [Fontes Java e composição](java-source-composition.json)',
       '- [Funções Java e linhas](java-methods.json)',
       '- [Fontes de integração e compilação](integration-build-source-index.json)',
       '- [Todas as entradas do APK](package-entries.json)',
       '- [Bibliotecas](native-libraries.json)',
       '- [DEX](dex-entries.json)','', '## Fontes consultadas','']
for s in all_sources:
    lines+=['- **'+str(s.get('id',''))+'** '+str(s.get('purpose') or s.get('description') or '')+' — '+str(s.get('path') or s.get('localPath') or '')+(' — [fonte]('+s['url']+')' if s.get('url') else '')]
args.markdown.parent.mkdir(parents=True,exist_ok=True);args.markdown.write_text('\n'.join(lines)+'\n',encoding='utf-8')
print(json.dumps({'sections':sum(len(g['sections']) for g in groups),'groups':len(groups),'document':str(args.docx),'markdown':str(args.markdown)}))
